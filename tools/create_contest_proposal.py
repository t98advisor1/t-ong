from __future__ import annotations

import os
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "submission"
TMP_DIR = Path(r"C:\tmp\t_ong_proposal_generated")
DOCX_PATH = OUT_DIR / "T-ong_공모전_제안서.docx"

SERVICE_URL = "https://t-ong.vercel.app"
VERSION = "v5.17.0-preview.35"

# Resolved document preset
# Base: grant_proposal (narrative_proposal)
# Named overrides: A4 compact Korean submission, Malgun Gothic, T-ong purple accent.
A4_WIDTH_DXA = 11906
PAGE_MARGIN_DXA = 864  # 0.6 in
CONTENT_WIDTH_DXA = A4_WIDTH_DXA - 2 * PAGE_MARGIN_DXA
FONT_NAME = "Malgun Gothic"
PURPLE = "6356C7"
DARK_PURPLE = "30285B"
BLUE = "2E74B5"
INK = "23233A"
MUTED = "666B7A"
PALE_PURPLE = "F3F1FF"
PALE_BLUE = "EFF5FF"
PALE_GRAY = "F5F6F8"
GRID = "D9DCE5"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, bold=None, color=INK, italic=None, name=FONT_NAME):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def set_table_borders(table, color=GRID, size=4):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_paragraph(paragraph, fill=PALE_PURPLE, left_border=PURPLE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), left_border)


def add_field(run, field_name: str):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_name
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_picture_alt(paragraph, description: str):
    drawings = paragraph._p.xpath(".//wp:docPr")
    if drawings:
        drawings[-1].set("descr", description)
        drawings[-1].set("title", description)


def add_title(doc: Document):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("기록은 한 번, 행정은 자동으로")
    set_run_font(r, 18, True, DARK_PURPLE)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(3)
    r = p2.add_run("관계형 디지털 교무수첩 T-ong")
    set_run_font(r, 15, True, PURPLE)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(7)
    r = p3.add_run("AI·디지털 기술을 활용한 2026년 학교 행정업무 경감 및 효율화 과제 공모전 제안서")
    set_run_font(r, 8.7, False, MUTED)
    r = p3.add_run(f"  |  작동 가능한 시제품  |  {VERSION}")
    set_run_font(r, 8.7, True, MUTED)


def add_heading(doc: Document, text: str, level=1):
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    return p


def add_body(doc: Document, text: str, label: str | None = None, after=2.5, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(9.5)
    if label:
        r = p.add_run(label)
        set_run_font(r, 9.5, True, DARK_PURPLE)
    r = p.add_run(text)
    set_run_font(r, 9.5, False, INK)
    return p


def add_compact_label(doc: Document, label: str, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.2)
    p.paragraph_format.line_spacing = 1.06
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(label)
    set_run_font(r, 9.35, True, PURPLE)
    r = p.add_run(text)
    set_run_font(r, 9.35, False, INK)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    set_run_font(r, 8.0, False, MUTED)
    return p


def get_image_font(size: int, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\malgunbd.ttf" if bold else r"C:\Windows\Fonts\malgun.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


def rounded_box(draw, xy, radius, fill, outline=None, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def make_workflow(path: Path):
    width, height = 2100, 520
    canvas = Image.new("RGB", (width, height), "FFFFFF")
    draw = ImageDraw.Draw(canvas)
    title_font = get_image_font(42, True)
    label_font = get_image_font(32, True)
    box_font = get_image_font(27, True)
    small_font = get_image_font(24, False)
    draw.text((50, 28), "업무 흐름 변화: 분산 기록에서 관계형 1회 입력으로", font=title_font, fill="#30285B")

    rows = [
        ("기존", "#8A8F9D", ["출결 확인", "별도 메모", "상담·관찰 재입력", "일정 별도 관리", "작성 때 다시 검색"]),
        ("T-ong", "#6356C7", ["수업 중 1회 입력", "학생번호·날짜로 연결", "관련 화면 즉시 조회", "내보내기·백업"]),
    ]
    y_positions = [145, 330]
    for (label, color, items), y in zip(rows, y_positions):
        rounded_box(draw, (45, y, 220, y + 110), 22, color)
        bbox = draw.textbbox((0, 0), label, font=label_font)
        tx = 132 - (bbox[2] - bbox[0]) / 2
        ty = y + 55 - (bbox[3] - bbox[1]) / 2 - 3
        draw.text((tx, ty), label, font=label_font, fill="FFFFFF")
        x = 260
        available = width - x - 45
        gap = 42
        box_w = int((available - gap * (len(items) - 1)) / len(items))
        for idx, item in enumerate(items):
            fill = "#F4F5F7" if label == "기존" else "#F3F1FF"
            outline = "#D8DBE2" if label == "기존" else "#B8B1F1"
            rounded_box(draw, (x, y, x + box_w, y + 110), 20, fill, outline, 3)
            lines = item.split("·") if "·" in item and len(item) > 11 else [item]
            if len(lines) == 1:
                bbox = draw.textbbox((0, 0), lines[0], font=box_font)
                draw.text((x + box_w / 2 - (bbox[2] - bbox[0]) / 2, y + 55 - (bbox[3] - bbox[1]) / 2 - 2), lines[0], font=box_font, fill=color)
            else:
                draw.multiline_text((x + box_w / 2, y + 29), "\n".join(lines), font=small_font, fill=color, anchor="ma", align="center", spacing=5)
            if idx < len(items) - 1:
                ax = x + box_w + 7
                ay = y + 55
                draw.line((ax, ay, ax + 26, ay), fill=color, width=6)
                draw.polygon([(ax + 26, ay), (ax + 14, ay - 9), (ax + 14, ay + 9)], fill=color)
            x += box_w + gap
    canvas.save(path, quality=95)


def fit_screenshot(img: Image.Image, box_w: int, box_h: int):
    img = img.convert("RGB")
    contained = ImageOps.contain(img, (box_w, box_h), Image.Resampling.LANCZOS)
    bg = Image.new("RGB", (box_w, box_h), "F7F8FB")
    x = (box_w - contained.width) // 2
    y = (box_h - contained.height) // 2
    bg.paste(contained, (x, y))
    return bg


def make_screenshot_grid(path: Path):
    width, height = 2100, 1340
    canvas = Image.new("RGB", (width, height), "FFFFFF")
    draw = ImageDraw.Draw(canvas)
    label_font = get_image_font(29, True)
    note_font = get_image_font(23, False)
    panels = [
        (ROOT / "tools" / "t-ong-preview11-home.png", "① 통합 홈", "자리·출결·일정·교실운영을 한 화면에서 이동"),
        (ROOT / "guide" / "20-attend.png", "② 출결 통합 입력", "상태·사유·메모 및 월별 내보내기"),
        (ROOT / "guide" / "20B-attend-student.png", "③ 관계형 학생 기록", "출결 화면에서 칭찬·관찰·세특·이력 연결"),
        (ROOT / "guide" / "22A-calendar-properties.png", "④ 일정·상담 연결", "사용자 속성과 상담 대상·후속 조치 기록"),
    ]
    margin = 35
    gap = 30
    panel_w = (width - margin * 2 - gap) // 2
    panel_h = (height - margin * 2 - gap) // 2
    for idx, (src, title, note) in enumerate(panels):
        col = idx % 2
        row = idx // 2
        x = margin + col * (panel_w + gap)
        y = margin + row * (panel_h + gap)
        rounded_box(draw, (x, y, x + panel_w, y + panel_h), 20, "FFFFFF", "#D9DCE5", 3)
        draw.rounded_rectangle((x + 3, y + 3, x + panel_w - 3, y + 76), radius=17, fill="#F3F1FF")
        draw.rectangle((x + 3, y + 45, x + panel_w - 3, y + 76), fill="#F3F1FF")
        draw.text((x + 24, y + 16), title, font=label_font, fill="#30285B")
        draw.text((x + 24, y + 82), note, font=note_font, fill="#5F6473")
        with Image.open(src) as image:
            fitted = fit_screenshot(image, panel_w - 28, panel_h - 150)
        canvas.paste(fitted, (x + 14, y + 134))
    canvas.save(path, quality=95)


def make_qr(path: Path):
    sys.path.insert(0, r"C:\tmp\tong_doc_deps")
    import qrcode

    qr = qrcode.QRCode(version=None, error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=14, border=3)
    qr.add_data(SERVICE_URL)
    qr.make(fit=True)
    image = qr.make_image(fill_color="#30285B", back_color="white").convert("RGB")
    image.save(path)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Inches(0.60)
    section.bottom_margin = Inches(0.60)
    section.left_margin = Inches(0.60)
    section.right_margin = Inches(0.60)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.30)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.10

    h1 = styles["Heading 1"]
    h1.font.name = FONT_NAME
    h1._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    h1.font.size = Pt(13)
    h1.font.bold = True
    h1.font.color.rgb = rgb(PURPLE)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = FONT_NAME
    h2._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    h2.font.size = Pt(10.5)
    h2.font.bold = True
    h2.font.color.rgb = rgb(DARK_PURPLE)
    h2.paragraph_format.space_before = Pt(4)
    h2.paragraph_format.space_after = Pt(2)
    h2.paragraph_format.keep_with_next = True

    caption = styles["Caption"]
    caption.font.name = FONT_NAME
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    caption.font.size = Pt(8)
    caption.font.color.rgb = rgb(MUTED)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hr = hp.add_run("2026 학교 행정업무 경감 및 효율화 과제 공모전  |  T-ong 교무수첩")
    set_run_font(hr, 7.2, False, MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fr = fp.add_run("T-ong 교무수첩  ·  ")
    set_run_font(fr, 7.2, False, MUTED)
    pr = fp.add_run()
    set_run_font(pr, 7.2, False, MUTED)
    add_field(pr, "PAGE")

    doc.core_properties.title = "기록은 한 번, 행정은 자동으로: 관계형 디지털 교무수첩 T-ong"
    doc.core_properties.subject = "AI·디지털 기술을 활용한 2026년 학교 행정업무 경감 및 효율화 과제 공모전 제안서"
    doc.core_properties.author = "T-ong 교무수첩"
    doc.core_properties.keywords = "학교 행정업무 경감, 디지털 교무수첩, 출결, 관계형 기록, 개인정보 보호"
    return doc


def add_page_one(doc: Document):
    add_title(doc)
    lead = doc.add_paragraph()
    lead.paragraph_format.space_before = Pt(2)
    lead.paragraph_format.space_after = Pt(6)
    lead.paragraph_format.left_indent = Pt(10)
    lead.paragraph_format.right_indent = Pt(8)
    lead.paragraph_format.line_spacing = 1.08
    shade_paragraph(lead)
    r = lead.add_run("핵심 제안  ")
    set_run_font(r, 9.6, True, PURPLE)
    r = lead.add_run("교사가 출결·상담·관찰·세특·학급 운영·일정을 각각 다시 기록하지 않도록, 학생번호와 날짜를 공통 기준으로 연결한다. 한 번 입력한 기록을 여러 업무 맥락에서 즉시 재사용하여 반복 입력과 누락 확인 부담을 줄이는 브라우저 기반 디지털 교무수첩이다.")
    set_run_font(r, 9.6, True, INK)

    add_heading(doc, "1. 현황(문제점)", 1)
    add_compact_label(doc, "가. 출결 확인의 반복성과 누락 부담  ", "담임·교과교사는 매시간 학생 상태를 확인하고 결석·지각·조퇴의 사유와 메모를 남긴다. 기록 시점이 늦어질수록 기억에 의존하거나 다시 확인해야 하며, 일별·월별 자료를 정리할 때 같은 내용을 재검토하는 부담이 커진다.")
    add_compact_label(doc, "나. 학생 기록의 분산  ", "출결, 상담, 관찰, 칭찬, 세부능력 및 특기사항, 자리 관계, 학급 운영 메모가 종이 수첩·스프레드시트·달력·개인 메모에 흩어진다. 학생 한 명의 맥락을 파악하려면 여러 자료를 다시 찾고 같은 사실을 반복 입력해야 한다.")
    add_compact_label(doc, "다. 현장 도구의 속도·접근 제약  ", "기존 Google Apps Script 기반 앱은 편리했으나 로그인 과정과 실행 속도에 제약이 있었다. 교실과 교과 이동 중에도 노트북·태블릿에서 즉시 열리고, 복잡한 설치나 계정 전환 없이 사용할 수 있는 도구가 필요하다.")
    add_compact_label(doc, "라. 편의성과 개인정보 보호의 동시 요구  ", "학생 자료를 중앙 서버에 모으면 접근성과 동기화는 좋아지지만 개인정보 유출·권한 관리 부담이 커진다. 반대로 로컬 기록만 사용하면 안전성은 높지만 업무 간 연결이 어렵다. 따라서 ‘로컬 우선 저장’과 ‘관계형 연결’을 함께 구현할 필요가 있다.")

    add_heading(doc, "2. 제안 내용", 1)
    add_heading(doc, "가. 제안 개요", 2)
    add_body(doc, "T-ong은 단순 자리배치 앱이 아니라, 학급과 학생을 중심으로 출결·상담·관찰·칭찬·세특·관계·이력을 연결하는 관계형 디지털 교무수첩이다. 학생 이름을 한 번 클릭하면 관련 기록을 같은 팝업에서 조회·추가할 수 있고, 일정의 상담 대상과 후속 조치도 학생 기록과 이어진다.")
    add_body(doc, "핵심은 ‘기록은 한 번, 활용은 여러 번’이다. 출결 화면에서 입력한 상태와 메모, 자리배치에서 입력한 관찰, 일정관리에서 등록한 상담 정보를 학생번호와 날짜 기준으로 연결하여 재검색·전사·중복 입력을 줄인다. 자료는 브라우저 localStorage에 저장하고 사용자가 백업·내보내기를 수행하도록 설계하였다.")

    add_heading(doc, "나. 설계 원칙", 2)
    add_compact_label(doc, "현장 즉시성  ", "순수 HTML·CSS·바닐라 JavaScript 기반으로 가볍게 구동하며, 노트북과 태블릿에서 수업 중 바로 입력할 수 있도록 화면을 최적화하였다.")
    add_compact_label(doc, "관계형 재사용  ", "학생번호·학급·날짜를 공통 키로 삼아 출결, 관찰, 칭찬, 세특, 상담, 자리 관계와 이력을 한 학생 맥락으로 묶는다.")
    add_compact_label(doc, "로컬 우선·사용자 통제  ", "학생·학급·출결·일정 자료는 현재 브라우저에만 저장하며, 외부 전송 없이 백업 파일과 엑셀·구글시트용 자료를 사용자가 직접 관리한다.")
    add_compact_label(doc, "AI의 보조적 활용  ", "AI를 학생 평가나 자동 의사결정에 사용하지 않고, ‘담임 5분 멘트’와 개발 과정의 코드 생성·검토 등 교사의 판단을 돕는 보조 수단으로 한정한다.")


def add_page_two(doc: Document, workflow: Path, screenshots: Path):
    doc.add_page_break()
    add_heading(doc, "2. 제안 내용(계속)", 1)
    add_heading(doc, "다. 주요 기능 구성", 2)
    add_compact_label(doc, "출결 통합 기록  ", "학급·날짜별로 출석·결석·지각·조퇴, 사유, 메모를 한 화면에서 입력하고 ‘오늘 전원 출석’, 월별 엑셀, 구글시트 복사 기능으로 후속 정리를 지원한다.")
    add_compact_label(doc, "학생 관계형 기록  ", "자리배치와 출결 화면에서 학생 이름을 누르면 종합·메모·관찰·세특·관계·이력 탭이 열린다. 칭찬은 +1과 -1을 모두 지원하고, 자리 고정·이동 및 가까이/멀리 배치 관계도 연결한다.")
    add_compact_label(doc, "일정·상담 관리  ", "월간·주간·오늘 일정을 제공하고 상담·출장·회의 및 사용자 정의 속성을 부여한다. 상담 대상 학생, 진행 상태, 후속 조치를 일정과 함께 기록하며 복사·내보내기를 지원한다.")
    add_compact_label(doc, "학급 운영 허브  ", "자리·모둠 배치, 1인 1역, 칭찬 포인트, 사다리타기, 룰렛, 타이머, 게임, 즐겨찾기, 도움말을 한 서비스에 통합하여 도구 전환 시간을 줄인다.")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(workflow), width=Inches(7.0))
    set_picture_alt(p, "기존 분산 기록과 T-ong 관계형 1회 입력 업무 흐름 비교")
    add_caption(doc, "[그림 1] 반복 입력·재검색 중심의 기존 흐름을 1회 입력·관계형 조회 흐름으로 전환")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(screenshots), width=Inches(6.95))
    set_picture_alt(p, "T-ong 통합 홈, 출결 입력, 학생 관계형 기록, 일정 상담 연결 화면")
    add_caption(doc, "[그림 2] 작동 가능한 시제품의 핵심 화면: 통합 홈·출결·학생 관계형 기록·일정/상담")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(1)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.left_indent = Pt(8)
    note.paragraph_format.right_indent = Pt(8)
    note.paragraph_format.line_spacing = 1.04
    shade_paragraph(note, fill=PALE_BLUE, left_border=BLUE)
    r = note.add_run("기술 구조  ")
    set_run_font(r, 8.8, True, BLUE)
    r = note.add_run("정적 웹앱(Vercel) → 교사 브라우저 localStorage에 업무 자료 저장 → 사용자가 백업·엑셀·구글시트로 내보내기. 서버에는 학생 자료를 수집하지 않으며, Upstash Redis에는 방문 통계용 익명 토큰과 횟수만 저장한다.")
    set_run_font(r, 8.8, False, INK)


def add_page_three(doc: Document, qr_path: Path):
    doc.add_page_break()
    add_heading(doc, "3. 추진 방법", 1)
    add_heading(doc, "가. 개발 및 전환", 2)
    add_body(doc, "기존 GAS 앱의 장점은 유지하되 로그인과 실행 속도 제약을 줄이기 위해 Vercel 기반 정적 웹앱으로 전환하였다. 기존 코드를 바탕으로 약 3일간 UI 전면 개편, 담임·교과교사용 기능 확장, 학생번호 기반 관계형 데이터 구조를 집중 개발하였다. 생성형 AI 코딩 도구와 코드 생성·검토·수정 사이클을 12회 이상 세션 단위로 반복하고, 실제 화면에서 기능을 검증하였다.", after=2.3)

    add_heading(doc, "나. 현장 적용 및 실증", 2)
    add_body(doc, "개발자인 현장 교사가 다음 학기 매시간 출결관리에 직접 사용한다. 적용 전 4주와 적용 후 4주를 비교하여 입력 완료 시간, 누락·정정 건수, 관련 기록 검색·취합 시간을 측정하고 사용성을 기록한다. 배포 전에는 원칙적으로 로컬·내부망에서 노트북과 태블릿 화면을 점검한 뒤 운영 주소에 반영한다.", after=2.3)

    add_heading(doc, "다. 확산 및 개선", 2)
    add_body(doc, "사용설명서 캐러셀과 화면별 도움말을 제공하고, 현장 피드백을 기능 단위로 반영한다. 별도 계정 없이 접속 가능한 무료 도구로 제공하되, 학교별 자료를 중앙 수집하지 않는 구조를 유지한다. 이후 실증 결과에 따라 출결 누락 알림, 기록 검색, 내보내기 양식 등 행정 경감 효과가 큰 기능부터 개선한다.", after=2.3)

    add_heading(doc, "라. 개인정보·보안 관리", 2)
    add_compact_label(doc, "학생 자료 비수집  ", "학생·학급·자리·출결·상담·관찰·세특·일정 자료는 브라우저 localStorage에만 저장하며 T-ong 서버나 Upstash Redis로 전송하지 않는다.")
    add_compact_label(doc, "통계 최소화  ", "Today는 당일 순 방문 브라우저 수, Total은 누적 접속 횟수이다. 브라우저의 무작위 식별자와 한국 날짜를 SHA-256으로 변환한 64자리 익명 토큰만 전송하며, 일별 집합은 3일 후 만료된다.")
    add_compact_label(doc, "이용자 통제와 한계 고지  ", "브라우저 삭제·기기 변경 시 자료가 사라질 수 있으므로 정기 백업을 안내한다. 공용 기기 사용 후에는 자료 삭제와 접근 통제를 권고하며, 통계에는 학생 이름·출결·상담 등 업무 자료가 포함되지 않는다.")

    add_heading(doc, "4. 기대효과", 1)
    add_body(doc, "아래 수치는 장기간 검증된 실적이 아니라 다음 학기 현장 적용에서 확인할 실증 목표이다. 동일 교사가 적용 전후의 업무 로그를 비교하여 과장 없이 검증한다.", after=3, first_indent=False)

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    headers = ["실증 지표", "목표", "측정 방법"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, PALE_PURPLE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, 8.3, True, DARK_PURPLE)
    rows = [
        ("일일 출결 입력 완료 시간", "50% 이상 단축", "적용 전·후 4주 평균 비교"),
        ("출결 누락·사후 정정 건수", "30% 이상 감소", "주별 누락·정정 로그 비교"),
        ("학생 관련 기록 검색·취합 시간", "40% 이상 단축", "상담·세특 준비 표본 업무 측정"),
        ("교사 사용성 만족도", "5점 중 4.0 이상", "신속성·편의성·신뢰성 자기평가"),
    ]
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 1 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_run_font(r, 8.05, idx == 1, PURPLE if idx == 1 else INK)
    widths = [3350, 1950, CONTENT_WIDTH_DXA - 5300]
    set_table_geometry(table, widths, indent_dxa=120)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])

    effect = doc.add_paragraph()
    effect.paragraph_format.space_before = Pt(4)
    effect.paragraph_format.space_after = Pt(2)
    effect.paragraph_format.line_spacing = 1.06
    r = effect.add_run("정성적 효과  ")
    set_run_font(r, 9.0, True, PURPLE)
    r = effect.add_run("기억과 재확인에 의존하던 업무를 기록 기반으로 전환하고, 담임·교과교사가 학생의 맥락을 빠르게 파악하도록 지원한다. 중앙 학생 DB를 만들지 않아 개인정보 관리 부담을 줄이면서도, 교사가 자신의 자료를 통제하는 현장 친화적 디지털 전환 모델을 제시한다.")
    set_run_font(r, 9.0, False, INK)

    qr_p = doc.add_paragraph()
    qr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr_p.paragraph_format.space_before = Pt(3)
    qr_p.paragraph_format.space_after = Pt(0)
    qr_p.add_run().add_picture(str(qr_path), width=Inches(0.72))
    set_picture_alt(qr_p, "T-ong 교무수첩 운영 주소 QR 코드")
    url_p = doc.add_paragraph()
    url_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    url_p.paragraph_format.space_before = Pt(0)
    url_p.paragraph_format.space_after = Pt(0)
    r = url_p.add_run(f"작동 시제품 확인  |  {SERVICE_URL}  |  {VERSION}")
    set_run_font(r, 8.5, True, DARK_PURPLE)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    workflow = TMP_DIR / "workflow.png"
    screenshots = TMP_DIR / "screenshots.png"
    qr_path = TMP_DIR / "t-ong-qr.png"
    make_workflow(workflow)
    make_screenshot_grid(screenshots)
    make_qr(qr_path)

    doc = setup_document()
    add_page_one(doc)
    add_page_two(doc, workflow, screenshots)
    add_page_three(doc, qr_path)
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
